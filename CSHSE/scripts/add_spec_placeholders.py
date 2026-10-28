#!/usr/bin/env python3
"""
Add per-SPECIFICATION placeholders to the prepared Reader Report templates.

The official form puts one Compliant / Non-Compliant box on each standard's
header row, but every specification row beneath it has its own (empty)
Compliant, Non-Compliant and Reader's Comments cells. Readers mark each
specification online, and lead readers send each reader's report to the
Council, so the download fills those per-spec cells instead of rolling
everything up to the standard (Lauri, Oct 2026).

Run AFTER prepare_reader_report_templates.py (it edits its output in place):
  python3 scripts/add_spec_placeholders.py

Adds, per level:
  {{c_<std>_<spec>}} {{n_<std>_<spec>}} {{cm_<std>_<spec>}} on the matched spec row
  {{reader_name}} {{report_date}} on the "Reader's Name and Credential" line
  {{reader_recommendation}} after "Indicate your recommendation to the Council"
and writes server/src/assets/reader-report-templates/spec-rows.json listing the
specs that got a row (any other spec's mark stays in the standard's comments).
Spec rows are matched to the level catalog by text (map_spec_rows.py).
"""
import copy, json, re, sys
from docx.shared import Pt
from map_spec_rows import mapping, uniq_cells, CAT

OUT = "server/src/assets/reader-report-templates"

def header_columns(row):
    """Positions (among distinct cells) of the header row's Compliant / Non-Compliant / Comments cells."""
    comp = non = comm = None
    for i, c in enumerate(uniq_cells(row)):
        t = re.sub(r"\{\{[^}]+\}\}", "", c.text).strip().lower()
        norm = re.sub(r"[^a-z]", "", t)
        if ("reader" in t and "comment" in t) or "areas of noncompliance" in t:
            comm = i if comm is None else comm
        elif norm.startswith("noncompliant"):
            non = i if non is None else non
        elif norm.endswith("compliant") and len(norm) <= 12:
            comp = i if comp is None else comp
    return comp, non, comm

def put(cell, token, size=9):
    p = cell.paragraphs[0]
    r = p.add_run(token)
    r.font.size = Pt(size)

def header_lookup(doc):
    """Map each table row to the nearest header row at or before it (document
    order, across tables) that names Compliant / Non-Compliant / Comments."""
    from map_spec_rows import tables
    out, last = {}, None
    for tbl in tables(doc):
        for row in tbl.rows:
            cols = header_columns(row)
            if all(i is not None for i in cols):
                last = (row, cols)
            out[id(row._tr)] = last
    return out

def fill_reader_line(doc):
    for p in doc.paragraphs:
        t = p.text
        if "Reader’s Name and Credential:" not in t and "Reader's Name and Credential:" not in t:
            continue
        if "{{reader_name}}" in t:
            return True
        if "{{inst_name}}" in t:
            # masters: one multi-label header paragraph. docx.patchDocument can't
            # fill a third/fourth token in that tab-split run, so drop the blank
            # "Reader's Name … Date:" labels from it and give the reader their
            # own line beneath.
            new = re.sub(r"\s*Reader[’']s Name and Credential:\s*Date:\s*$", "", t)
            first = p.runs[0]
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            run = p.add_run(new)
            run.bold, run.italic = first.bold, first.italic
            if first.font.size: run.font.size = first.font.size
            from docx.text.paragraph import Paragraph
            line = copy.deepcopy(p._p)
            for r in list(line.iterchildren()):
                if r.tag.endswith("}r"): line.remove(r)
            p._p.addnext(line)
            run2 = Paragraph(line, p._parent).add_run("Reader’s Name and Credential: {{reader_name}}     Date: {{report_date}}")
            run2.bold = first.bold
            if first.font.size: run2.font.size = first.font.size
            return True
        first = p.runs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        run = p.add_run("Reader’s Name and Credential: {{reader_name}}     Date: {{report_date}}")
        run.bold = first.bold
        if first.font.size: run.font.size = first.font.size
        return True
    return False

def add_recommendation(doc):
    for p in doc.paragraphs:
        if "Indicate your recommendation to the Council" in p.text:
            if "{{reader_recommendation}}" in doc.element.xml:
                return True
            new_p = copy.deepcopy(p._p)
            for r in list(new_p.iterchildren()):
                if r.tag.endswith("}r"):
                    new_p.remove(r)
            p._p.addnext(new_p)
            from docx.text.paragraph import Paragraph
            Paragraph(new_p, p._parent).add_run("{{reader_recommendation}}")
            return True
    return False

def process(level):
    doc, _secs, res = mapping(level)
    placed, skipped = [], []
    headers = header_lookup(doc)
    for (std, spec), ((sc, _i), row, _crit) in res.items():
        if row is None or sc < 0.5:
            skipped.append(f"{std}.{spec}"); continue
        hdr = headers.get(id(row._tr))
        comp, non, comm = hdr[1] if hdr else (None, None, None)
        # The decision columns are the RIGHTMOST distinct cells, so count from
        # the right: headers and rows differ in their leading/merged cells
        # (curriculum rows have an extra blank first cell; masters Standard 12's
        # header spans an extra grid column).
        cells = uniq_cells(row)
        width = len(uniq_cells(hdr[0])) if hdr else 0
        comp, non, comm = (len(cells) - (width - i) if i is not None else None for i in (comp, non, comm))
        blank = lambda c: not re.sub(r"[\W_]", "", c.text)
        ok = hdr is not None and all(i is not None and 0 < i < len(cells) and blank(cells[i])
                                     for i in (comp, non, comm))
        if not ok:
            skipped.append(f"{std}.{spec}"); continue
        key = f"{std}_{spec}"
        put(cells[comp], f"{{{{c_{key}}}}}", 11)
        put(cells[non], f"{{{{n_{key}}}}}", 11)
        put(cells[comm], f"{{{{cm_{key}}}}}", 8)
        placed.append(f"{std}.{spec}")
    rl = fill_reader_line(doc)
    rec = add_recommendation(doc)
    doc.save(f"{OUT}/{level}.docx")
    print(f"  {level}: {len(placed)} spec rows, skipped {skipped}, reader line={rl}, recommendation={rec}")
    return placed

if __name__ == "__main__":
    manifest = {}
    for level in ("associate", "baccalaureate", "masters"):
        manifest[level] = process(level)
    json.dump(manifest, open(f"{OUT}/spec-rows.json", "w"), indent=1)
    total = sum(len(CAT[l][s]["specs"]) for l in CAT for s in CAT[l])
    print(f"placed {sum(map(len, manifest.values()))} of {total} catalog specs")
