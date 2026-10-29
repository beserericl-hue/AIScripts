#!/usr/bin/env python3
"""
Give every numbered PART of a specification, and every Introduction row, its
own Compliant / Non-Compliant boxes on the Reader Report templates.

The official forms list numbered parts under many specifications (4.a has
1. Measurable student learning outcomes, 2. Assessment plan, …) and each one is
marked by the reader; the Introduction lists one row per item (degree offered,
the institution, …). The download marked only the specification / one roll-up
Introduction box (Lauri, Oct 2026).

Run AFTER add_spec_placeholders.py (edits the templates in place):
  python3 scripts/add_part_placeholders.py

Adds, per level:
  {{c_<std>_<spec>_p<n>}} {{n_<std>_<spec>_p<n>}}   on the n-th part row of a spec
  {{c_introduction_<code>}} {{n_…}} {{cm_…}}          on the Introduction row for an
                                                      INTRO_RUBRIC item (matched by text)
  {{c_<intro>_p<n>}} {{n_…}}                          on any other Introduction row (the
                                                      bacc Program Information rows follow
                                                      the certification check)
  {{cert_included}}                                   in the certification "Included: Yes/No" cell
and writes server/src/assets/reader-report-templates/spec-parts.json.
"""
import json, re, subprocess
from docx.shared import Pt
from map_spec_rows import uniq_cells, tables, words
from add_spec_placeholders import header_columns, put, OUT

# Rows inside a spec block that are not parts of it.
NOT_A_PART = re.compile(r"reader.?s noted program strengths|^matrix\b|^general program characteristics", re.I)


def intro_rubric():
    js = ("import { INTRO_RUBRIC } from './server/src/data/introRubric';"
          "console.log(JSON.stringify(INTRO_RUBRIC.map(r => ({ code: r.specCode, text: r.title + ' ' + r.criteria }))));")
    out = subprocess.run(["npx", "--prefix", "server", "tsx", "-e", js], capture_output=True, text=True, check=True).stdout
    return json.loads(out.strip().splitlines()[-1])


LIST_MARK = re.compile(r"^\s*(\d+|[a-z]|[ivx]+)[.)]\s*")


def unmarked(label):
    """The label without its list marker ("2.", "b.", "iv)")."""
    return LIST_MARK.sub("", label)


def shaded(cell):
    fills = re.findall(r'w:fill="([0-9A-Fa-f]{6})"', cell._tc.xml)
    return any(f.upper() not in ("FFFFFF",) for f in fills)


def overlap(a, b):
    """Share of the row's words found in the rubric text; 0 unless ≥2 shared
    (so a short label like "Institution Name" doesn't match by one word)."""
    A, B = set(words(a)[:20]), set(words(b)[:30])
    return len(A & B) / max(1, min(len(A), len(B))) if len(A & B) >= 2 else 0


def box_columns(row, hdr):
    """Distinct-cell positions of (compliant, non-compliant, comments) for a row,
    counted from the right like add_spec_placeholders."""
    if not hdr:
        return None
    cells = uniq_cells(row)
    width = len(uniq_cells(hdr[0]))
    cols = tuple(len(cells) - (width - i) if i is not None else None for i in hdr[1])
    if not all(i is not None and 0 < i < len(cells) for i in cols):
        return None
    return cols


def process(level, rubric):
    from docx import Document
    doc = Document(f"{OUT}/{level}.docx")
    # Track the nearest header row as we go (document order, across tables).
    # (Keying rows by id(row._tr) is unreliable: lxml reuses proxy ids.)
    parts, intro, parent, n, hdr = [], [], None, 0, None
    for tbl in tables(doc):
        for row in tbl.rows:
            cols = header_columns(row)
            if all(i is not None for i in cols):
                hdr = (row, cols)
            cells = uniq_cells(row)
            text = " ".join(dict.fromkeys(c.text.strip() for c in cells if c.text.strip()))
            # Certification gate: "Included: Yes/No"
            for c in cells:
                if c.text.strip().startswith("Included: Yes/No") and "{{cert_included}}" not in c.text:
                    put(c, "  {{cert_included}}", 11)
                    intro.append({"key": "cert_included", "intro": "cert"})
            m = re.search(r"\{\{c_((?:intro_[a-f])|introduction|\d+(?:_[a-z]+)?)\}\}", text)
            if m:
                parent, n = m.group(1), 0
                continue
            if len(cells) == 1 or parent is None:
                if len(cells) == 1 and text:  # a heading row ends the block (a blank spacer doesn't)
                    parent = None
                continue
            is_intro = parent == "introduction" or parent.startswith("intro_")
            if not is_intro and "_" not in parent:
                continue  # the line under a standard header, not a spec part
            cols = box_columns(row, hdr)
            if cols is None:
                continue
            comp, non, comm = cols
            blank = lambda c: not re.sub(r"[\W_]", "", c.text)
            # The row's text is its last non-empty cell before the boxes (some rows
            # carry a stray letter cell, e.g. "B", in front).
            label = next((c.text.strip() for c in reversed(cells[:comp]) if c.text.strip()), "")
            if (not label or not all(blank(cells[i]) for i in cols) or shaded(cells[0])
                    or NOT_A_PART.search(label) or unmarked(label)[:1].islower()):
                continue
            # The baccalaureate form's "Program Information" block (intro_a:
            # Certification Page uploaded, Program Name, …) is the certification
            # check, not the A. Introduction items — those rows follow it.
            if is_intro and parent != "intro_a":
                best = max(rubric, key=lambda r: overlap(label, r["text"]))
                if best["code"] != "cert" and overlap(label, best["text"]) >= 0.5:
                    key = f"introduction_{best['code']}"
                    put(cells[comp], f"{{{{c_{key}}}}}", 11)
                    put(cells[non], f"{{{{n_{key}}}}}", 11)
                    if not any(e.get("key") == key for e in intro):
                        put(cells[comm], f"{{{{cm_{key}}}}}", 8)
                    intro.append({"key": key, "intro": best["code"], "text": label})
                    continue
            n += 1
            key = f"{parent}_p{n}"
            put(cells[comp], f"{{{{c_{key}}}}}", 11)
            put(cells[non], f"{{{{n_{key}}}}}", 11)
            entry = {"key": key, "n": n, "text": unmarked(label)}
            if is_intro:
                entry["introGroup"] = parent
                if parent == "intro_a":
                    entry["follows"] = "cert"
                intro.append(entry)
            else:
                std, spec = parent.split("_", 1)
                entry["spec"] = f"{std}.{spec}"
                parts.append(entry)
    doc.save(f"{OUT}/{level}.docx")
    print(f"  {level}: {len(parts)} spec parts, {len(intro)} introduction rows")
    return {"parts": parts, "intro": intro}


if __name__ == "__main__":
    rubric = intro_rubric()
    manifest = {level: process(level, rubric) for level in ("associate", "baccalaureate", "masters")}
    json.dump(manifest, open(f"{OUT}/spec-parts.json", "w"), indent=1, ensure_ascii=False)
